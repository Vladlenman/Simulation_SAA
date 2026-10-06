"""Extract the data the BM tool needs out of a BM-Tool Excel workbook.

The workbook's ``timeseries`` sheet is a hard copy of ``BB_timeseries`` (the
Bloomberg pull), so everything this script produces works without a Bloomberg
Terminal.  Layout of ``timeseries``: column pairs, odd column = observation
date + series name in row 1, even column = monthly total return in PERCENT.

Writes
    data/monthly_returns.csv   month (YYYY-MM) x ticker, returns as decimals
    data/benchmarks.json       catalogue: name, currency, first/last month, n
    data/reference_figures.json  the figures BM-Tool itself reported, for tests

Usage
    python tools/import_excel.py path/to/BM-Tool.xlsx
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required for the import: pip install openpyxl")

from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

# BM-Tool!C4:C18 -> the benchmark each SAA row was mapped to, and the SAA
# weights that were live in the delivered workbook.  Used to build the
# reference scenario and the regression test.
EXCEL_SAA_ROWS = [
    ("Anleihen EURO", "LBEATREU Index", 0.29),
    ("Anleihen HTM / Geldmarkt", "ESTR3MA Index (+ 100BP)", 0.35),
    ("Anleihen EURO (1-3y)", "LET1TREU Index", 0.00),
    ("Anleihen EURO (Overnight)", "DBDCONIA Index", 0.00),
    ("Anleihen Welt", "LEGATRUU Index", 0.03),
    ("Anleihen EM", "JPEIDHEU Index", 0.00),
    ("Wandelanleihen (EUR hedged)", "Convertibles Index EUR", 0.04),
    ("Aktien Welt", "MSDEWIN Index", 0.08),
    ("Aktien Welt (EUR hedged)", "MXWOHEUR Index", 0.00),
    ("Aktien Europa (EUR hedged)", "NDDLE15 Index", 0.05),
    ("Aktien EM", "MSDEEEMN Index", 0.02),
    ("Alternative Investments", "ESTR3MA Index (+ 100BP)", 0.03),
    ("Alternative Investments (CTA)", "NEIXCTA Index", 0.00),
    ("Mikrofinanz", "ESTR3MA Index (+ 100BP)", 0.02),
    ("Immobilien", "ESTR3MA Index (+ 200BP)", 0.09),
]

# BM-Tool!D34:I34 for "SAA neu" over its own start date, with TER 1.09 % p.a.
# Sharpe is deliberately absent: the workbook computes it against an excess
# return column that is only filled from row 123 (2011-01) onwards, so its
# 0.8285 is a hybrid of gross and excess returns.  See docs/.
EXCEL_REFERENCE_FIGURES = {
    "scenario": "SAA neu (BM-Tool Abgabe)",
    "ter": 0.0109,
    "start": "2003-02",
    "end": "2025-12",
    "months": 275,
    "total_return": 0.8629210691663325,
    "cagr": 0.02752004633094418,
    "volatility": 0.0284192838044391,
    "max_drawdown": -0.09305372257332678,
    "calmar": 0.2957436367928027,
    "excel_sharpe_do_not_trust": 0.8284639760907965,
}


def month_key(value: dt.datetime) -> str:
    return f"{value.year:04d}-{value.month:02d}"


def read_timeseries(path: Path) -> tuple[dict[str, dict[str, float]], dict[str, dict]]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book["timeseries"].iter_rows(values_only=True))
    header = rows[0]

    meta_by_name = read_basicdata(book)

    series: dict[str, dict[str, float]] = {}
    catalogue: dict[str, dict] = {}
    for col in range(0, len(header), 2):
        raw_name = header[col]
        if raw_name is None or not str(raw_name).strip():
            continue
        name = str(raw_name).strip()

        observations: dict[str, float] = {}
        for row in rows[1:]:
            if col + 1 >= len(row):
                continue
            stamp, value = row[col], row[col + 1]
            if not isinstance(stamp, dt.datetime):
                continue
            if value is None or isinstance(value, str):
                continue
            # Excel stores these as percent; the tool works in decimals.
            observations[month_key(stamp)] = float(value) / 100.0

        if not observations:
            # e.g. JPCAEU1Y / JGAGGUSD: a header with a broken Bloomberg pull.
            continue

        months = sorted(observations)
        series[name] = observations
        meta = meta_by_name.get(name, {})
        catalogue[name] = {
            "name": name,
            "description": meta.get("description", ""),
            "currency": meta.get("currency", ""),
            "source_column": get_column_letter(col + 1),
            "first_month": months[0],
            "last_month": months[-1],
            "months": len(months),
        }

    book.close()
    return series, catalogue


def read_basicdata(book) -> dict[str, dict]:
    """Ticker -> currency / long name, from the ``basicdata`` sheet."""
    out: dict[str, dict] = {}
    if "basicdata" not in book.sheetnames:
        return out
    for row in book["basicdata"].iter_rows(min_row=2, max_col=5, values_only=True):
        ticker = row[0]
        if not ticker or not str(ticker).strip():
            continue
        currency = row[3] if len(row) > 3 else None
        description = row[4] if len(row) > 4 else None

        def clean(value) -> str:
            text = "" if value is None else str(value).strip()
            return "" if text.startswith("#") else text

        out[str(ticker).strip()] = {
            "currency": clean(currency),
            "description": clean(description),
        }
    return out


def write_csv(series: dict[str, dict[str, float]], path: Path) -> None:
    tickers = sorted(series)
    months = sorted({m for obs in series.values() for m in obs})
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["month", *tickers])
        for month in months:
            writer.writerow(
                [month]
                + [
                    "" if month not in series[t] else repr(series[t][month])
                    for t in tickers
                ]
            )


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__)
        return 2
    source = Path(argv[1]).expanduser()
    if not source.exists():
        sys.exit(f"workbook not found: {source}")

    series, catalogue = read_timeseries(source)
    write_csv(series, DATA / "monthly_returns.csv")
    (DATA / "benchmarks.json").write_text(
        json.dumps(catalogue, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    (DATA / "reference_figures.json").write_text(
        json.dumps(
            {
                "source_workbook": source.name,
                "imported": dt.date.today().isoformat(),
                "excel_saa_rows": EXCEL_SAA_ROWS,
                "figures": EXCEL_REFERENCE_FIGURES,
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )

    print(f"{len(series)} series -> {DATA / 'monthly_returns.csv'}")
    for name, info in sorted(catalogue.items()):
        print(
            f"  {name:32s} {info['months']:4d} obs  "
            f"{info['first_month']} .. {info['last_month']}  {info['currency']}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))

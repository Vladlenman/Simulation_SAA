"""Build ``data/returns.xlsx`` - the single source of return data.

The workbook is what you fill from Bloomberg: one sheet, one column per
series, one row per month. Everything the tool charts comes from here, the
money-market and book-yield series included, so nothing is computed behind
your back.

    python tools/build_workbook.py                     rebuild from current data
    python tools/build_workbook.py --from BM-Tool.xlsx import the Excel first

Values are stored **in percent**, the way Bloomberg's
``DAY_TO_DAY_TOT_RETURN_GROSS_DVDS`` delivers them and the way the old
``timeseries`` sheet held them, so a BDH pull can be pasted in unchanged.
The ``Info`` sheet records that unit; set it to ``Dezimal`` there if you would
rather keep decimals.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl wird gebraucht:  pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

SHEET = "Renditen (in Prozent)"
INFO_SHEET = "Info"
HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")


def build(destination: Path) -> Path:
    """Rewrite the workbook from the current data, derived columns included."""
    from saa.data import panel

    data = panel()  # market series plus everything the tool derives
    return write_workbook(
        destination,
        months=list(data.months),
        tickers=sorted(data.tickers),
        get_column=data.column,
        describe=data.describe,
    )


def write_workbook(destination, months, tickers, get_column, describe) -> Path:
    """Write ``returns.xlsx``: one sheet of returns, one sheet of provenance."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = SHEET

    sheet.cell(row=1, column=1, value="Monat")
    for index, ticker in enumerate(tickers, start=2):
        sheet.cell(row=1, column=index, value=ticker)

    series = {t: get_column(t) for t in tickers}
    for row, month in enumerate(months, start=2):
        sheet.cell(row=row, column=1, value=month)
        for index, ticker in enumerate(tickers, start=2):
            value = series[ticker][row - 2]
            if value == value:  # not NaN
                cell = sheet.cell(row=row, column=index, value=float(value) * 100.0)
                cell.number_format = "0.0000"

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 10
    for index in range(2, len(tickers) + 2):
        sheet.column_dimensions[get_column_letter(index)].width = 15

    info = book.create_sheet(INFO_SHEET)
    rows: list[tuple] = [
        ("Einheit", "Prozent"),
        ("", "Prozent = 1,25 bedeutet 1,25 % Monatsrendite."),
        ("", "Dezimal = 0,0125 fuer dieselbe Rendite."),
        ("", "Nur dieses Feld entscheidet, wie die Zahlen gelesen werden."),
        ("", ""),
        ("Erstellt am", dt.date.today().isoformat()),
        ("Monate", f"{months[0]} bis {months[-1]}"),
        ("Zeitreihen", len(tickers)),
        ("", ""),
        ("Zeitreihe", "Herkunft"),
    ]
    for ticker in tickers:
        meta = describe(ticker)
        origin = {
            "market": "Marktdaten",
            "derived": "berechnet",
            "schedule": "Verzinsungspfad",
            "example": "BEISPIEL - bitte ersetzen",
        }.get(meta.get("origin", "market"), meta.get("origin", ""))
        note = meta.get("description") or ""
        rows.append((ticker, f"{origin}{f' - {note}' if note else ''}"))

    for index, (left, right) in enumerate(rows, start=1):
        info.cell(row=index, column=1, value=left).font = Font(bold=index in (1, 10))
        info.cell(row=index, column=2, value=right)
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 70

    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)

    # Keep the catalogue in step with the columns actually written, so every
    # series - derived ones included - still states its origin and coverage.
    catalogue: dict[str, dict] = {}
    for ticker in tickers:
        meta = dict(describe(ticker))
        known = [m for m, v in zip(months, series[ticker]) if v == v]
        meta.update(
            {
                "name": ticker,
                "first_month": known[0] if known else None,
                "last_month": known[-1] if known else None,
                "months": len(known),
            }
        )
        catalogue[ticker] = meta
    (destination.parent / "benchmarks.json").write_text(
        json.dumps(catalogue, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return destination


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--from",
        dest="source",
        metavar="BM-TOOL.XLSX",
        help="zuerst aus einer BM-Tool-Datei importieren",
    )
    parser.add_argument(
        "--out",
        default=str(ROOT / "data" / "returns.xlsx"),
        help="Zieldatei (Standard: data/returns.xlsx)",
    )
    args = parser.parse_args(argv[1:])

    if args.source:
        from tools.import_excel import main as import_main

        if import_main(["import_excel", args.source]) != 0:
            return 1
        from saa.data import reset_caches

        reset_caches()

    written = build(Path(args.out))
    from saa.data import panel, reset_caches

    reset_caches()
    data = panel()
    print(f"{written}")
    print(f"  {len(data.tickers)} Zeitreihen, {len(data.months)} Monate "
          f"({data.months[0]} bis {data.months[-1]}), Werte in Prozent")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
